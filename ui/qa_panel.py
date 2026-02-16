from pathlib import Path

import streamlit as st
from docx import Document as DocxDocument

from RAG.generation import GenerationError, generate_answer
from RAG.ingestion.loaders import process_uploads_for_session
from RAG.retrieval import RetrievalError, retrieve_context
from state.config import EMBEDDING_DIMENSIONS


# -------------------------
# Controls
# -------------------------

DEFAULT_MODEL_OPTIONS = [
    "gpt-4.1-2025-04-14",
    "gpt-4o-2024-08-06",
    "gpt-5",
    "gpt-5-thinking",
    "gpt-5-pro",
    "o4-mini-2025-04-16",
    "o4-mini-deep-research-2025-06-26",
]


def _is_allowed_model(model_id: str) -> bool:
    """Check whether a model ID is allowed in the selector."""
    mid = (model_id or "").lower()
    return mid.startswith("o") or (mid.startswith("gpt-") and "codex" not in mid)


def load_model_options(client):
    """Load and cache selectable model IDs."""
    cached = st.session_state.get("available_models")
    if cached:
        return cached

    if client is None:
        return DEFAULT_MODEL_OPTIONS

    try:
        data = client.models.list()
        models = [m.id for m in getattr(data, "data", [])]
        options = sorted({m for m in models if _is_allowed_model(m)})
        if options:
            st.session_state.available_models = options
            return options
    except Exception:
        pass

    return DEFAULT_MODEL_OPTIONS

def model_selector(options):
    """Select LLM for answer generation."""
    current = st.session_state.get("gpt_model")
    index = options.index(current) if current in options else 0
    st.session_state.gpt_model = st.selectbox(
        "🤖 Select model:",
        options,
        index=index,
    )


def advanced_controls():
    """Diversity, GraphRAG, and creativity controls."""
    with st.expander("🎛️ Advanced Controls"):
        st.toggle(
            "Use GraphRAG retrieval",
            key="use_graphrag",
            help="Include knowledge-graph context in retrieval when a GraphRAG index is loaded.",
        )
        
        st.slider(
            "Diversity", 0.0, 1.0, 0.7, 0.01, key="diversity"
        )

        if st.session_state.gpt_model in {
            "gpt-4o-2024-08-06",
            "gpt-4.1-2025-04-14",
        }:
            st.slider(
                "Creativity", 0.0, 1.0, 0.3, 0.01, key="temperature"
            )
        else:
            st.session_state.temperature = 0.3


# -------------------------
# Ask (Knowledge Base) a Question UI
# -------------------------
def qa_panel(client):
    st.header("❓ Ask a Question")

    has_key = st.session_state.get("api_verified", False)
    has_kb = st.session_state.get("index") is not None
    disabled = not (has_key and has_kb)

    if not has_key:
        st.info("ℹ️ Please set a valid API key.")
    elif not has_kb:
        st.info("ℹ️ Please build or load a Knowledge-Base.")

    # ---- defaults
    st.session_state.setdefault("gpt_model", "gpt-4.1-2025-04-14")
    st.session_state.setdefault("last_query", "")
    st.session_state.setdefault("last_answer", "")
    st.session_state.setdefault("context_meta", "")
    st.session_state.setdefault("use_graphrag", False)

    model_selector(load_model_options(client))
    advanced_controls()

    query = st.text_area(
        "Ask your question:",
        height=280,
        disabled=disabled,
    )

    st.markdown("#### 📎 Add any relevant file(s) for this question (optional)")
    _uploaded_files = st.file_uploader(
        "Upload PDFs/TXT/CSV/XLSX or images:",
        type=[
            "pdf", "txt", "csv", "xlsx",
            "png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff",
        ],
        accept_multiple_files=True,
        disabled=disabled,
    )
    
    col1, _, col3 = st.columns([1, 2, 1])
    answer_clicked = col1.button(
        "💬 Answer", use_container_width=True, disabled=disabled
    )
    save_clicked = col3.button(
        "💾 Save last Q&A", use_container_width=True, disabled=disabled
    )

    if answer_clicked and query:
        index_path = st.session_state.get("index_path")
        meta_path = st.session_state.get("meta_path")
        if not index_path or not meta_path:
            st.error("Missing index/meta file paths in session. Reload or rebuild the KB.")
            return

        with st.spinner("📂 Processing uploaded files..."):
            try:
                n_text, n_imgs = process_uploads_for_session(
                    uploaded_files=_uploaded_files,
                    client=client,
                    embedding_model=st.session_state.embedding_model,
                    dimension=EMBEDDING_DIMENSIONS[st.session_state.embedding_model],
                )
                if n_text or n_imgs:
                    st.success(
                        f"✅ Processed {n_text} text chunks and {n_imgs} image(s) from the uploaded files!"
                    )
            except Exception as exc:
                st.session_state.upload_db = None
                st.session_state.upload_meta = []
                st.session_state.upload_images = []
                st.error(f"❌ Upload processing failed: {exc}")

            if st.session_state.get("upload_meta") or st.session_state.get("upload_images"):
                st.caption(
                    f"Uploads ready: {len(st.session_state.get('upload_meta', []))} text chunks, "
                    f"{len(st.session_state.get('upload_images', []))} images!"
                )

        with st.spinner("🔍 Retrieving context..."):
            try:
                retrieval_result = retrieve_context(
                    query=query,
                    index_path=index_path,
                    meta_path=meta_path,
                    graphrag_dir=st.session_state.get("graphrag"),
                    api_key=st.session_state.get("api_key"),
                    diversity=st.session_state.get("diversity", 0.7),
                    use_graphrag=st.session_state.use_graphrag,
                )
                context = retrieval_result["context_text"]
            except RetrievalError as exc:
                st.error(f"❌ Retrieval failed: {exc}")
                return

            upload_meta = st.session_state.get("upload_meta") or []
            if upload_meta:
                upload_lines = [
                    f"[{m['source']} | chunk {m['chunk_id']}]: {m['text']}"
                    for m in upload_meta
                ]
                context = context + "\n\n" + "\n\n".join(upload_lines)
            if st.session_state.get("upload_images"):
                for img in st.session_state.upload_images:
                    context += f"\n\n[uploaded/{img['name']} | image]: (image attached)"
        st.markdown("### 💡 Answer")

        system = (
            "You are an expert scientific research assistant. Use the context provided from research papers "
            "to answer the user query accurately and with source-aware details. If context is insufficient, "
            "state that clearly before providing a best-effort response."
        )
        try:
            answer = generate_answer(
                query=query,
                context_text=context,
                model=st.session_state.gpt_model,
                system=system,
                api_key=st.session_state.get("api_key"),
                enable_web_search=False,
            )
        except GenerationError as exc:
            st.error(f"❌ Generation failed: {exc}")
            return

        st.session_state.update(
            last_query=query,
            last_answer=answer,
            context_meta=context,
        )

    if save_clicked and st.session_state.last_answer:
        path = Path("outputs/qa_pairs.docx")
        doc = DocxDocument(path) if path.exists() else DocxDocument()
        doc.add_heading("Question:", 2)
        doc.add_paragraph(st.session_state.last_query)
        doc.add_heading("Answer:", 2)
        doc.add_paragraph(st.session_state.last_answer)
        path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(path)
        st.success(f"✅ Saved to {path}")

    if st.session_state.context_meta:
        st.divider()
        with st.expander("📚 Retrieved Context"):
            st.markdown(
                st.session_state.context_meta.replace("\n", "<br>"),
                unsafe_allow_html=True,
            )
