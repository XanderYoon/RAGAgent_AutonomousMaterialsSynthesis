from pathlib import Path
import re
import fitz
import tiktoken
from openai import OpenAI
import numpy as np
import faiss
import streamlit as st
import pickle
import os
import openai
from langchain.text_splitter import RecursiveCharacterTextSplitter
import pycountry
import time
from langchain_community.vectorstores import FAISS 
from langchain_community.docstore.in_memory import InMemoryDocstore
from langchain_core.documents import Document
from langchain_openai import OpenAIEmbeddings, ChatOpenAI
import pandas as pd
import io
import base64
import mimetypes
from docx import Document as DocxDocument
from langchain_community.retrievers import BM25Retriever
from langchain.retrievers import EnsembleRetriever, MultiQueryRetriever
from langchain.retrievers.contextual_compression import ContextualCompressionRetriever
from langchain.retrievers.document_compressors import LLMChainExtractor
from langchain.retrievers.document_compressors import CrossEncoderReranker
from langchain_community.cross_encoders import HuggingFaceCrossEncoder
from langchain_community.document_transformers import LongContextReorder
from langchain_experimental.graph_transformers import LLMGraphTransformer
import pickle
import networkx as nx
from langchain_core.retrievers import BaseRetriever
from langchain.globals import set_llm_cache
from langchain_community.cache import SQLiteCache
import json
import subprocess
import yaml



# s3: Q&A

if "index" in st.session_state:
    for key in ["last_query", "last_answer", "context_meta"]:
        if key not in st.session_state:
            st.session_state[key] = ""    
    query = st.text_area("Ask your question here:", height=280, placeholder="Type your question...")
    
    # Uploads UI kept from the extended app
    st.markdown("#### 📎 Add any relevant file(s) for this question (optional)")
    uploaded_files = st.file_uploader(
        "Upload PDFs/TXT/CSV/XLSX or images (PNG/JPG/JPEG/GIF/BMP/TIF/TIFF):",
        type=["pdf", "txt", "csv", "xlsx", "png", "jpg", "jpeg", "gif", "bmp", "tif", "tiff"],
        accept_multiple_files=True
    )
    use_uploads = True

    ## Answer
    # if st.button("💬 Answer") and query:
    col1, col2, col3 = st.columns([1, 2, 1])     
    with col1:
        answer_clicked = st.button("💬 Answer", use_container_width=True)
    with col3:
        save_clicked = st.button("💾 Save last Q&A", use_container_width=True)        
    st.markdown("")        
    if answer_clicked and query:
    ##  
        # 1) process uploads (text + images), cache for reuse
        with st.spinner("📂 Processing uploaded files..."):        
            if uploaded_files:
                try:
                    up_db, up_meta, up_images = build_upload_bundle(
                        uploaded_files=uploaded_files,
                        client=client,
                        embedding_model=st.session_state.embedding_model,
                        dimension=d_em2dim[st.session_state.embedding_model]
                    )
                    st.session_state.upload_db = up_db
                    st.session_state.upload_meta = up_meta
                    st.session_state.upload_images = up_images
                    n_text = len(up_meta)
                    n_imgs = len(up_images)
                    st.success(f"✅ Processed {n_text} text chunks and {n_imgs} image(s) from the uploaded files!")
                except Exception as e:
                    st.error(f"❌ Upload processing failed: {e}")
            else:
                st.session_state.upload_db = None
                st.session_state.upload_meta = []
                st.session_state.upload_images = []
    
            if use_uploads and "upload_meta" in st.session_state:
                st.caption(f"Uploads ready: {len(st.session_state.get('upload_meta', []))} text chunks, "
                           f"{len(st.session_state.get('upload_images', []))} images!")

        # 2) retrieve context (KB + optional uploads) via MMR
        with st.spinner("🔍 Knowledge-Base: Retrieving relevant context ..."):        
            query_embedding = client.embeddings.create(
                input=query,
                model=st.session_state.embedding_model
            ).data[0].embedding
            query_embedding = np.array(query_embedding, dtype="float32").reshape(1, -1)
            flat_emb = query_embedding[0].astype(float).tolist()
            ##
            ##
            # results = st.session_state.db.max_marginal_relevance_search_by_vector(
            #     embedding=flat_emb, k=top_k_textKBfaiss, fetch_k=top_k_textKBfaiss * 2, lambda_mult=1.0 - diversity
            # )
            vs_retriever = st.session_state.db.as_retriever(
                search_type="mmr",
                search_kwargs={
                    "k": top_k_textKBfaiss,
                    "fetch_k": top_k_textKBfaiss * 2,
                    "lambda_mult": 1.0 - diversity,
                },
            )
            bm25 = st.session_state.bm25
            hybrid = EnsembleRetriever(retrievers=[vs_retriever, bm25], weights=[0.5, 0.5])      
            llm_expander = ChatOpenAI(model="gpt-4o-mini")
            multi_query = MultiQueryRetriever.from_llm(
                retriever=hybrid,
                llm=llm_expander,
                include_original=True,
            ) 
            # cross_encoder = HuggingFaceCrossEncoder(model_name="cross-encoder/ms-marco-MiniLM-L-6-v2")
            # reranker = CrossEncoderReranker(model=cross_encoder, top_n=int(top_k_textKBfaiss+top_k_textKBbm))        
            # reranked_exapanded = ContextualCompressionRetriever(
            #     base_retriever=multi_query,
            #     base_compressor=reranker,
            # )        
            compressor_llm = ChatOpenAI(model="gpt-4o-mini", temperature=0)  
            compressor = LLMChainExtractor.from_llm(compressor_llm)    
            # compressed_retriever = ContextualCompressionRetriever(
            #     base_retriever=reranked_exapanded,
            #     base_compressor=compressor,
            # )        
            compressed_retriever = ContextualCompressionRetriever(
                base_retriever=multi_query,
                base_compressor=compressor,
            )                    
            results = compressed_retriever.invoke(query)
            reorder = LongContextReorder()
            results = reorder.transform_documents(results) 
            ##
            ## ==> till now (KB-text-faiss, KB-text-bm, KB-text-Engi)
            ## 
            results_up = []
            if use_uploads and st.session_state.get("upload_db") is not None:
                ##
                # results_up = st.session_state.upload_db.max_marginal_relevance_search_by_vector(
                #     embedding=flat_emb, k=top_k_addKBfaiss, fetch_k=top_k_addKBfaiss * 2, lambda_mult=1.0 - diversity
                # )
                #
                vs_retriever_up = st.session_state.upload_db.as_retriever(
                    search_type="mmr",
                    search_kwargs={
                        "k": top_k_addKBfaiss,
                        "fetch_k": top_k_addKBfaiss * 2,
                        "lambda_mult": 1.0 - diversity,
                    },
                )
                multi_query_up = MultiQueryRetriever.from_llm(
                    retriever=vs_retriever_up,
                    llm=llm_expander,
                    include_original=True,
                )
                compressed_retriever_up = ContextualCompressionRetriever(
                    base_retriever=multi_query_up,
                    base_compressor=compressor,
                )
                results_up = compressed_retriever_up.invoke(query)
                results_up = reorder.transform_documents(results_up)                
                ##
            merged = results        
            if results_up:
                merged.extend(results_up)
            context_meta_chunks = [
                f"[{doc.metadata['source']} | chunk {doc.metadata['chunk_id']}]: {doc.page_content}"
                for doc in merged
            ]        
            original_cmc = [
                f"[{doc.metadata['source']} | chunk {doc.metadata['chunk_id']}]: {doc.metadata['original_content']}"
                for doc in merged
            ]    
            ## ==> till now (upload-text-faiss, upload-text-Engi)
            ##
            if use_uploads and st.session_state.get("upload_images"):
                for img in st.session_state.upload_images[:]:
                    context_meta_chunks.append(f"[uploaded/{img['name']} | image]: (image attached)")  
                    original_cmc.append(f"[uploaded/{img['name']} | image]: (image attached)")  
            ## ==> till now (upload-image)
            ##
        ##
        with st.spinner("Knowledge-Graph: Retrieving entities, relationships, and summaries..."):        
            graph_context_chunks = []
            graph_context = ""
            original_gc = ""
            if "graphrag" in st.session_state and os.path.isdir(st.session_state.graphrag):
                try:
                    ROOT_DIR = Path(st.session_state.graphrag)
                    result = subprocess.run(
                        [
                            "graphrag", "query",
                            "--root", str(ROOT_DIR),
                            "--method", "global",
                            "--query", query
                        ],
                        capture_output=True,
                        text=True
                    )
                    if result.returncode !=0:
                        raise Exception(f"⚠️ Knowledge-Graph query failed:\n{result.stderr.strip()}")
                    answer = result.stdout.strip()
                    if answer:
                        # graph_context = (
                        #     "\n[GraphRAG Context]:\n" +
                        #     textwrap.shorten(answer, width=6000, placeholder=" ...")
                        # )
                        graph_context = (
                            "\n[Knowledge-Graph Context]:\n" + answer
                        )                        
                        context_meta_chunks.append(graph_context)
                        original_cmc.append(graph_context)
                        st.success("✅ Knowledge-Graph context successfully retrieved and added!")
                    else:
                        st.info(f"⚠️ No Knowledge-Graph response found for this query!")
                except Exception as e:
                    st.warning(f"⚠️ Error while retrieving from Knowledge-Graph: {e}")
            else:
                st.info("ℹ️ No Knowledge-Graph directory loaded - skipping graph-based retrieval.")
            ## ==> till now (KB-text-graph) 
        ##
        context_meta = "\n\n".join(context_meta_chunks)
        original_cm = "\n\n".join(original_cmc)
            
        # 3) build prompt & call model
        system_instructions = (
            "You are an expert scientific research assistant. Use the context provided from research papers to answer "
            "the user query as accurately as possible. Provide detailed responses using as much of the provided context "
            "as possible. If the answer is not clearly found in the context, respond with: "
            "'The context does not provide enough information to answer this question.' and default to your parametric knowledge to give a response "
            "If the provided context is not enough to answer the user, use web search to find relevant information. "
            "At the end of your answer, also mention the source file names or uploaded image names referenced in the context "
            "(e.g., [DL-rheed-harris-SI.pdf | chunk 1]) or (e.g., [uploaded/image.png | image])."
        )

        prompt_text = f"""
User query: {query}

--- BEGIN CONTEXT ---
{context_meta}
--- END CONTEXT ---

Answer:
""".strip()

        # Prepare image content blocks (if supported by model route)
        supports_images = st.session_state.gpt_model in ["gpt-4o-2024-08-06", "gpt-4.1-2025-04-14"]
        user_content = [{"type": "text", "text": prompt_text}]
        if use_uploads and st.session_state.get("upload_images") and supports_images:
            for img in st.session_state.upload_images[:]:
                user_content.append({"type": "text", "text": f"[uploaded/{img['name']} | image]:"})
                user_content.append({"type": "image_url", "image_url": {"url": img["data_url"]}})

        # Route per model
        model_choice = st.session_state.gpt_model

        try:
##
            st.markdown("### 💡 Answer")
            placeholder = st.empty()
            answer_text = ""
            
            if model_choice == "o4-mini-deep-research-2025-06-26":
                with client.responses.stream(
                    model=model_choice,
                    instructions=system_instructions,
                    tools=[{"type": "web_search_preview"}],
                    input=prompt_text,
                ) as stream:
                    for event in stream:
                        if event.type == "response.output_text.delta":
                            answer_text += event.delta
                            placeholder.markdown(answer_text + "▌")
                            time.sleep(stream_delay)  
                        elif event.type == "response.error":
                            st.error(str(event.error))
                        elif event.type in {"response.output_text.done", "response.completed"}:
                            placeholder.markdown(answer_text)     
                            
            elif model_choice in ["gpt-4o-2024-08-06", "gpt-4.1-2025-04-14"]:
                with client.chat.completions.stream(
                    model=model_choice,
                    messages=[
                        {"role": "system", "content": system_instructions},
                        {"role": "user", "content": user_content},
                    ],
                    temperature=temperature,
                ) as stream:
                    for event in stream:
                        if event.type == "content.delta":
                            answer_text += event.delta
                            placeholder.markdown(answer_text + "▌")
                            time.sleep(stream_delay)
                        elif event.type == "content.done":
                            placeholder.markdown(answer_text)
            
            elif model_choice in {"gpt-5", "gpt-5-thinking", "gpt-5-pro"}:
                model_id = "gpt-5" if model_choice == "gpt-5-thinking" else model_choice
                kwargs = {"model": model_id, "instructions": system_instructions, "input": prompt_text}
                if model_choice == "gpt-5-thinking":
                    kwargs["reasoning"] = {"effort": "high"}
                try:
                    with client.responses.stream(**kwargs) as stream:
                        for event in stream:
                            if event.type == "response.output_text.delta":
                                answer_text += event.delta
                                placeholder.markdown(answer_text + "▌")
                                time.sleep(stream_delay)
                            elif event.type == "response.error":
                                st.error(str(event.error))
                            elif event.type in {"response.output_text.done", "response.completed"}:
                                placeholder.markdown(answer_text)
                except Exception as e:
                    st.warning(f"⚠️ Streaming not supported for this model: {e}")
                    response = client.responses.create(**kwargs)
                    answer_text = response.output_text
                    placeholder.markdown(answer_text)            
            else:
                with client.chat.completions.stream(
                    model=model_choice,
                    messages=[
                        {"role": "system", "content": system_instructions},
                        {"role": "user", "content": prompt_text},
                    ],
                ) as stream:
                    for event in stream:
                        if event.type == "content.delta":
                            answer_text += event.delta
                            placeholder.markdown(answer_text + "▌")
                            time.sleep(stream_delay)
                        elif event.type == "content.done":
                            placeholder.markdown(answer_text)
## 
            # with st.expander("📚 Retrieved Context for this Query"):
            #     context_meta_html = original_cm.replace("\n", "<br>")
            #     # context_meta_html = context_meta.replace("\n", "<br>")                
            #     st.markdown(f"<div style='overflow-wrap: break-word; width: 600px'>{context_meta_html}</div>", unsafe_allow_html=True)

           ##
            st.session_state.last_query = query
            st.session_state.last_answer = answer_text
            st.session_state.context_meta = original_cm          
            # st.session_state.context_meta = context_meta                    
            ##            

        except Exception as e:
            st.error(f"❌ Inference failed: {e}")

    ##
    ## save
    if save_clicked and st.session_state.last_answer:
        try:
            save_path = Path("outputs/qa_pairs.docx")
            if save_path.exists():
                doc = DocxDocument(save_path)
            else:
                doc = DocxDocument()
                doc.add_heading("Question–Answer Log", level=1)
                doc.add_paragraph("")

            doc.add_heading("Question:", level=2)
            doc.add_paragraph(st.session_state.last_query.strip())
            doc.add_heading("Answer:", level=2)
            doc.add_paragraph(st.session_state.last_answer.strip())
            doc.add_paragraph("")  # spacing
            save_path.parent.mkdir(parents=True, exist_ok=True)
            doc.save(save_path)
            st.success(f"✅ Q&A pair added to: {save_path}")
        except Exception as e:
            st.error(f"❌ Saving Q&A pair failed: {e}")
    ##

    ## retrieved context
    if st.session_state.context_meta:
        st.divider()
        with st.expander("📚 Retrieved Context for this Query"):
            context_meta_html = st.session_state.context_meta.replace("\n", "<br>")            
            st.markdown(f"<div style='overflow-wrap: break-word; width: 600px'>{context_meta_html}</div>", unsafe_allow_html=True)
    ##   

else:
    st.info("⚠️ Please build or load a Knowledge-Base before asking a question!")